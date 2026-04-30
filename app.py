import os
import io
import base64
import subprocess
import tempfile
import shutil
from datetime import datetime
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from openpyxl import load_workbook
import anthropic

app = Flask(__name__)
CORS(app)

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "Formulario_de_Solicitacao_de_Acesso.xlsx")
ANTHROPIC_KEY = os.environ.get("ANTHROPIC_API_KEY", "")

# ─────────────────────────────────────────
# ROTA: Extrair dados via Claude Vision
# ─────────────────────────────────────────
@app.route("/extrair-cliente", methods=["POST"])
def extrair_cliente():
    try:
        client = anthropic.Anthropic(api_key=ANTHROPIC_KEY)
        content = []

        for key in ["rg", "conta"]:
            f = request.files.get(key)
            if f:
                data = base64.standard_b64encode(f.read()).decode("utf-8")
                media_type = f.mimetype or "image/jpeg"
                if media_type == "application/pdf":
                    content.append({
                        "type": "document",
                        "source": {"type": "base64", "media_type": "application/pdf", "data": data}
                    })
                else:
                    content.append({
                        "type": "image",
                        "source": {"type": "base64", "media_type": media_type, "data": data}
                    })

        content.append({
            "type": "text",
            "text": (
                "Analise os documentos enviados (documento de identidade e/ou conta de energia elétrica) "
                "e extraia os dados. Retorne SOMENTE um JSON válido, sem markdown, sem explicações:\n\n"
                '{"nome":"","cpf":"","rg":"","data_expedicao":"","endereco":"","cep":"",'
                '"municipio":"","uf":"","fixo":"","conta_contrato":"","classe":"",'
                '"tensao":"","carga":"","ramo":"","disjuntor":""}\n\n'
                "Se algum campo não estiver disponível retorne string vazia."
            )
        })

        resp = client.messages.create(
            model="claude-sonnet-4-20250514",
            max_tokens=1000,
            messages=[{"role": "user", "content": content}]
        )
        text = resp.content[0].text.strip()
        text = text.replace("```json", "").replace("```", "").strip()
        import json
        data = json.loads(text)
        return jsonify({"ok": True, "data": data})

    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500


# ─────────────────────────────────────────
# ROTA: Extrair equipamentos via Claude Vision
# ─────────────────────────────────────────
@app.route("/extrair-equipamentos", methods=["POST"])
def extrair_equipamentos():
    try:
        client = anthropic.Anthropic(api_key=ANTHROPIC_KEY)
        content = []

        for key in ["eq", "inv"]:
            f = request.files.get(key)
            if f:
                data = base64.standard_b64encode(f.read()).decode("utf-8")
                media_type = f.mimetype or "image/jpeg"
                if media_type == "application/pdf":
                    content.append({
                        "type": "document",
                        "source": {"type": "base64", "media_type": "application/pdf", "data": data}
                    })
                else:
                    content.append({
                        "type": "image",
                        "source": {"type": "base64", "media_type": media_type, "data": data}
                    })

        content.append({
            "type": "text",
            "text": (
                "Analise os documentos de módulos fotovoltaicos e/ou inversores solares. "
                "Retorne SOMENTE um JSON válido, sem markdown:\n\n"
                '{"modulos":[{"potencia_w":null,"quantidade":null,"fabricante":"","modelo":""}],'
                '"inversores":[{"fabricante":"","modelo":"","potencia_kw":null,"tensao":"",'
                '"corrente_a":null,"fator_potencia":null,"rendimento_pct":null,"dht":""}]}'
            )
        })

        resp = client.messages.create(
            model="claude-sonnet-4-20250514",
            max_tokens=1000,
            messages=[{"role": "user", "content": content}]
        )
        text = resp.content[0].text.strip()
        text = text.replace("```json", "").replace("```", "").strip()
        import json
        data = json.loads(text)
        return jsonify({"ok": True, "data": data})

    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500


# ─────────────────────────────────────────
# ROTA: Gerar PDF preenchendo o Excel original
# ─────────────────────────────────────────
@app.route("/gerar-pdf", methods=["POST"])
def gerar_pdf():
    try:
        import json
        dados = request.get_json()
        cliente = dados.get("cliente", {})
        equipamentos = dados.get("equipamentos", {})

        # Criar cópia temporária do template
        tmp_dir = tempfile.mkdtemp()
        xlsx_path = os.path.join(tmp_dir, "formulario.xlsx")
        shutil.copy(TEMPLATE_PATH, xlsx_path)

        wb = load_workbook(xlsx_path)

        # ── ABA 1: Dados do Cliente ──
        ws1 = wb["1"]

        def s(val):
            return str(val).strip() if val else ""

        # Nome
        ws1["C10"] = s(cliente.get("nome"))
        # CPF/CNPJ
        ws1["R10"] = s(cliente.get("cpf"))
        # RG
        ws1["AC9"] = s(cliente.get("rg"))
        # Data expedição
        exp = s(cliente.get("data_expedicao"))
        if exp:
            try:
                ws1["AC10"] = datetime.strptime(exp, "%Y-%m-%d")
            except:
                ws1["AC10"] = exp
        # Endereço
        ws1["C13"] = s(cliente.get("endereco"))
        # CEP
        ws1["D15"] = s(cliente.get("cep"))
        # Município
        ws1["I15"] = s(cliente.get("municipio"))
        # UF
        ws1["Q15"] = s(cliente.get("uf"))
        # Celular (sempre fixo)
        ws1["T13"] = "(86)9 8128-1094"
        # Fixo
        ws1["Z13"] = s(cliente.get("fixo"))
        # Email (sempre fixo)
        ws1["V15"] = "contato@electrize.com.br"
        # Conta Contrato
        ws1["Z17"] = s(cliente.get("conta_contrato"))
        # Tipo de Solicitação
        ws1["G19"] = s(cliente.get("tipo_solicitacao")) or "CONEXÃO DE GD EM UNIDADE CONSUMIDORA EXISTENTE SEM AUMENTO DE POTÊNCIA DISPONIBILIZADA (ver item abaixo)"
        # Ramo de Atividade
        ws1["G25"] = s(cliente.get("ramo"))
        # Classe
        ws1["F27"] = s(cliente.get("classe")) or "Residencial"
        # Tipo de Ligação
        ws1["T27"] = s(cliente.get("ligacao"))
        # Tensão
        tensao = cliente.get("tensao")
        if tensao:
            try:
                ws1["AC27"] = int(tensao)
            except:
                ws1["AC27"] = tensao
        # Carga Declarada
        carga = cliente.get("carga")
        if carga:
            try:
                ws1["F29"] = float(str(carga).replace(",", "."))
            except:
                pass
        # Disjuntor
        disj = cliente.get("disjuntor")
        if disj:
            try:
                ws1["P29"] = int(disj)
            except:
                pass
        # Tipo de Ramal
        ws1["F31"] = s(cliente.get("ramal")) or "AÉREO"
        # Coordenadas UTM
        ws1["P33"] = s(cliente.get("utm_x"))
        ws1["Y33"] = s(cliente.get("utm_y"))
        # Modalidade
        ws1["I53"] = s(cliente.get("modalidade")) or "AUTOCONSUMO LOCAL"
        # Data início de operação
        dt_op = s(cliente.get("data_operacao"))
        if dt_op:
            ws1["U59"] = dt_op

        # ── ABA 0: Equipamentos ──
        ws0 = wb["0"]

        modulos = equipamentos.get("modulos", [])
        for i, m in enumerate(modulos[:10]):
            row = 7 + i
            if m.get("potencia_w"):
                try:
                    ws0.cell(row=row, column=4).value = float(str(m["potencia_w"]).replace(",", "."))
                except:
                    pass
            if m.get("quantidade"):
                try:
                    ws0.cell(row=row, column=8).value = int(m["quantidade"])
                except:
                    pass
            if m.get("fabricante"):
                ws0.cell(row=row, column=20).value = m["fabricante"]
            if m.get("modelo"):
                ws0.cell(row=row, column=27).value = m["modelo"]

        inversores = equipamentos.get("inversores", [])
        for i, inv in enumerate(inversores[:30]):
            row = 22 + i
            if inv.get("fabricante"):
                ws0.cell(row=row, column=4).value = inv["fabricante"]
            if inv.get("modelo"):
                ws0.cell(row=row, column=8).value = inv["modelo"]
            if inv.get("potencia_kw"):
                try:
                    ws0.cell(row=row, column=12).value = float(str(inv["potencia_kw"]).replace(",", "."))
                except:
                    pass
            if inv.get("tensao"):
                ws0.cell(row=row, column=16).value = inv["tensao"]
            if inv.get("corrente_a"):
                try:
                    ws0.cell(row=row, column=20).value = float(str(inv["corrente_a"]).replace(",", "."))
                except:
                    pass
            if inv.get("fator_potencia"):
                try:
                    ws0.cell(row=row, column=23).value = float(str(inv["fator_potencia"]).replace(",", "."))
                except:
                    pass
            if inv.get("rendimento_pct"):
                try:
                    ws0.cell(row=row, column=26).value = float(str(inv["rendimento_pct"]).replace(",", "."))
                except:
                    pass
            if inv.get("dht"):
                ws0.cell(row=row, column=29).value = inv["dht"]

        wb.save(xlsx_path)

        # Converter Excel → PDF via LibreOffice
        result = subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf",
             "--outdir", tmp_dir, xlsx_path],
            capture_output=True, text=True, timeout=60
        )

        pdf_path = xlsx_path.replace(".xlsx", ".pdf")

        if not os.path.exists(pdf_path):
            return jsonify({
                "ok": False,
                "error": "LibreOffice falhou ao converter",
                "stdout": result.stdout,
                "stderr": result.stderr
            }), 500

        with open(pdf_path, "rb") as f:
            pdf_data = f.read()

        shutil.rmtree(tmp_dir, ignore_errors=True)

        return send_file(
            io.BytesIO(pdf_data),
            mimetype="application/pdf",
            as_attachment=True,
            download_name="Formulario_Microgeracao.pdf"
        )

    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500


# ─────────────────────────────────────────
# ROTA: Health check
# ─────────────────────────────────────────
@app.route("/", methods=["GET"])
def health():
    return jsonify({"status": "ok", "service": "SolarForm API"})


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
